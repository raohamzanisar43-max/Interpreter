"""Excel export for admins (openpyxl)."""
import io
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from compliance.models import AuditLog
from organisations.models import Company
from rating.models import AnonymousItem, ChecklistScore, PairwiseTask
from rounds.models import KeyPoint, Round
from submissions.models import Assignment

from .scoring import company_aggregates, item_stats, r1, rank_companies


def _sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(str(h)) + 4))
    ws.freeze_panes = "A2"
    return ws


def build_workbook(rounds):
    wb = Workbook()
    wb.remove(wb.active)
    rounds = list(rounds)
    _sheet(wb, "Rounds", ["ID", "Name", "Language", "Deadline", "Status", "Raters/item", "Pairwise/item", "Retention days"],
           [[r.pk, r.name, r.language, r.deadline.replace(tzinfo=None), r.status, r.min_raters_per_item,
             r.pairwise_per_item, r.retention_days] for r in rounds])
    _sheet(wb, "Companies", ["ID", "Name", "Contact", "Status", "Terms version", "Terms accepted at"],
           [[c.pk, c.name, c.contact_email, c.status, c.terms_version_accepted,
             c.terms_accepted_at.replace(tzinfo=None) if c.terms_accepted_at else None] for c in Company.objects.all()])
    rows = []
    for a in Assignment.objects.filter(round__in=rounds).select_related("round", "interpreter__company"):
        taken = a.interpretations.filter(deleted_at=None).count()
        rows.append([a.round.name, a.interpreter.company.name, a.interpreter.full_name, a.interpreter.email, a.status,
                     a.submitted_at.replace(tzinfo=None) if a.submitted_at else None,
                     f"{taken}/{a.round.recordings.count()}"])
    _sheet(wb, "Interpreters", ["Round", "Company", "Interpreter", "Email", "Status", "Submitted at", "Recordings"], rows)
    rows = [[k.source_recording.round.name, k.source_recording.title, k.order + 1, k.text, k.category, k.weight, k.admin_notes]
            for k in KeyPoint.objects.filter(source_recording__round__in=rounds).select_related("source_recording__round")]
    _sheet(wb, "Recordings & Key points", ["Round", "Recording", "#", "Key point", "Category", "Weight", "Admin notes"], rows)
    rows = []
    for s in ChecklistScore.objects.filter(rating_task__anonymous_item__round__in=rounds).select_related(
            "key_point", "rating_task__rater__user", "rating_task__anonymous_item__round",
            "rating_task__anonymous_item__interpretation__assignment__interpreter__company",
            "rating_task__anonymous_item__interpretation__source_recording"):
        t = s.rating_task
        it = t.anonymous_item
        rows.append([it.round.name, it.public_code, it.interpretation.assignment.interpreter.company.name,
                     it.interpretation.source_recording.title, t.rater.user.email, s.key_point.text, s.value, s.comment,
                     t.status, "yes" if t.excluded else ""])
    _sheet(wb, "Raw checklist scores", ["Round", "Item", "Company", "Recording", "Rater", "Key point", "Value", "Comment",
                                        "Task status", "Excluded"], rows)
    rows = [[p.round.name, p.source_recording.title, p.rater.user.email, p.item_a.public_code,
             p.item_a.interpretation.assignment.interpreter.company.name, p.item_b.public_code,
             p.item_b.interpretation.assignment.interpreter.company.name, p.winner, p.comment]
            for p in PairwiseTask.objects.filter(round__in=rounds).select_related(
                "round", "source_recording", "rater__user", "item_a__interpretation__assignment__interpreter__company",
                "item_b__interpretation__assignment__interpreter__company")]
    _sheet(wb, "Pairwise", ["Round", "Recording", "Rater", "A", "A company", "B", "B company", "Winner", "Comment"], rows)
    rows = []
    for r in rounds:
        aggs = company_aggregates(r)
        ranks, _ = rank_companies(aggs)
        for pk, a in aggs.items():
            rows.append([r.name, a["name"], r1(a["score"]), ranks.get(pk), a["items"], r1(a["pairwise"]),
                         "; ".join(f"{c['category']}: {r1(c['score'])}" for c in a["per_category"])])
    _sheet(wb, "Company aggregates", ["Round", "Company", "Score", "Rank", "Items", "Pairwise win rate", "By category"], rows)
    counts = Counter(AuditLog.objects.values_list("action", flat=True))
    _sheet(wb, "Audit summary", ["Action", "Count"], [[k, v] for k, v in sorted(counts.items())])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
